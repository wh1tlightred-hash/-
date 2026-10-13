"""报告导出：Markdown → HTML（目录 + 图片内嵌为单文件）与 Word（.docx）。

- Markdown：原样落盘，图片为相对路径（便于在仓库里浏览，验收脚本也按相对路径核对）。
- HTML：自动生成目录；图片默认以 base64 内嵌，单个文件即可发给他人打开。
- Word：用 python-docx 把同一份 Markdown 转成 .docx（标题、段落、粗体/代码、列表、表格、图片、
  引用），并插入一个目录域——在 Word 中右键「更新域」即可生成页码目录。
"""
from __future__ import annotations

import base64
import re
from pathlib import Path

import markdown as md_lib
from markdown.extensions.toc import slugify_unicode

_HTML_TEMPLATE = """<!DOCTYPE html>
<html lang="zh-CN">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>中医声诊三分类实验报告</title>
<style>
:root{{--ink:#0b0b0b;--ink2:#52514e;--muted:#898781;--line:#e1e0d9;--accent:#1c5cab;--wash:#f0efec;--bg:#fcfcfb;}}
body{{font-family:"Microsoft YaHei","PingFang SC","SimHei",sans-serif;max-width:1040px;margin:24px auto;padding:0 20px;line-height:1.75;color:var(--ink);background:var(--bg);}}
h1{{border-bottom:3px solid var(--accent);padding-bottom:8px;}}
h2{{border-bottom:1px solid var(--line);padding-bottom:4px;margin-top:40px;color:var(--accent);}}
h3{{margin-top:28px;}}
table{{border-collapse:collapse;display:block;overflow-x:auto;max-width:100%;margin:12px 0;font-size:13.5px;}}
th,td{{border:1px solid var(--line);padding:5px 8px;text-align:center;white-space:nowrap;}}
th{{background:var(--wash);}}
img{{max-width:100%;border:1px solid var(--line);margin:8px 0;background:#fff;}}
code{{background:var(--wash);padding:1px 5px;border-radius:3px;font-size:92%;}}
pre{{background:var(--wash);padding:10px;border-radius:5px;overflow-x:auto;}}
blockquote{{border-left:4px solid var(--accent);margin:12px 0;padding:4px 16px;color:var(--ink2);background:#f7fafc;}}
.toc{{background:var(--wash);border-radius:6px;padding:8px 20px;margin:16px 0 24px;font-size:14px;}}
.toc ul{{margin:2px 0;padding-left:20px;}}
.toc a{{color:var(--accent);text-decoration:none;}}
@media print{{.toc{{page-break-after:always;}} img{{page-break-inside:avoid;}}}}
</style>
</head>
<body>
{body}
</body>
</html>
"""

_IMG_MD = re.compile(r"^!\[(?P<alt>[^\]]*)\]\((?P<src>[^)]+)\)\s*$")
_MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".svg": "image/svg+xml"}


def _insert_toc_marker(report_md: str) -> str:
    """在第一个二级标题前插入 [TOC]，目录只在 HTML 中展开（Markdown 原文不受影响）。"""
    idx = report_md.find("\n## ")
    if idx < 0:
        return report_md
    return report_md[:idx] + "\n\n**目录**\n\n[TOC]\n" + report_md[idx:]


def _embed_images(html: str, base_dir: Path) -> str:
    """把 <img src="相对路径"> 替换为 data URI，生成可单独分发的 HTML。找不到的图片保持原样。"""
    def repl(m):
        src = m.group(2)
        if src.startswith(("data:", "http://", "https://")):
            return m.group(0)
        p = (base_dir / src).resolve()
        if not p.exists():
            return m.group(0)
        mime = _MIME.get(p.suffix.lower(), "application/octet-stream")
        data = base64.b64encode(p.read_bytes()).decode("ascii")
        return f'{m.group(1)}data:{mime};base64,{data}{m.group(3)}'

    return re.sub(r'(<img[^>]*?src=")([^"]+)(")', repl, html)


def markdown_to_html(report_md: str, base_dir: Path, embed_images: bool = True) -> str:
    body = md_lib.markdown(
        _insert_toc_marker(report_md),
        extensions=["tables", "fenced_code", "sane_lists", "toc"],
        extension_configs={"toc": {"slugify": slugify_unicode, "toc_depth": "2-3"}},
    )
    if embed_images:
        body = _embed_images(body, base_dir)
    return _HTML_TEMPLATE.format(body=body)


# ------------------------------------------------------------------ Word

_INLINE = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\*[^*\s][^*]*\*)")


def _set_cjk_font(style, name: str = "Microsoft YaHei"):
    from docx.oxml.ns import qn

    style.font.name = name
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = rpr.makeelement(qn("w:rFonts"), {})
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), name)


def _add_inline(paragraph, text: str):
    """把 **粗体**、`代码`、*斜体* 转成 run。"""
    from docx.shared import RGBColor

    for part in _INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.color.rgb = RGBColor(0x1C, 0x5C, 0xAB)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2:
            paragraph.add_run(part[1:-1]).italic = True
        else:
            paragraph.add_run(part)


def _add_toc_field(doc):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    p = doc.add_paragraph()
    run = p.add_run()
    for tag, attrs, text in (("w:fldChar", {"w:fldCharType": "begin"}, None),
                             ("w:instrText", {"xml:space": "preserve"}, 'TOC \\o "1-3" \\h \\z \\u'),
                             ("w:fldChar", {"w:fldCharType": "separate"}, None)):
        el = OxmlElement(tag)
        for k, v in attrs.items():
            el.set(qn(k), v)
        if text:
            el.text = text
        run._r.append(el)
    p.add_run("（在 Word 中右键此处 →「更新域」即可生成目录）").italic = True
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    p.add_run()._r.append(end)


def _add_table(doc, rows: list[list[str]]):
    from docx.shared import Pt

    header, body = rows[0], rows[1:]
    ncols = len(header)
    table = doc.add_table(rows=1, cols=ncols)
    table.style = "Light Grid Accent 1"
    for j, cell in enumerate(header):
        table.rows[0].cells[j].text = ""
        _add_inline(table.rows[0].cells[j].paragraphs[0], cell)
    for r in body:
        cells = table.add_row().cells
        for j in range(ncols):
            cells[j].text = ""
            _add_inline(cells[j].paragraphs[0], r[j] if j < len(r) else "")
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(8 if ncols > 6 else 9)
    doc.add_paragraph()


def _split_row(line: str) -> list[str]:
    return [c.strip() for c in line.strip().strip("|").split("|")]


def markdown_to_docx(report_md: str, base_dir: Path, out_path: Path) -> None:
    """把报告 Markdown 转为 Word。支持本报告用到的 Markdown 子集。"""
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm, Pt, RGBColor

    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin = sec.bottom_margin = Cm(2.0)
    for name in ("Normal", "Title", "Heading 1", "Heading 2", "Heading 3", "Quote",
                 "List Bullet", "List Number", "Caption"):
        try:
            _set_cjk_font(doc.styles[name])
        except KeyError:
            pass
    doc.styles["Normal"].font.size = Pt(10.5)
    for name, color in (("Heading 1", (0x1C, 0x5C, 0xAB)), ("Heading 2", (0x0B, 0x0B, 0x0B))):
        doc.styles[name].font.color.rgb = RGBColor(*color)

    lines = report_md.splitlines()
    i, toc_done = 0, False
    page_width = sec.page_width - sec.left_margin - sec.right_margin
    while i < len(lines):
        line = lines[i]
        s = line.strip()
        if not s or s == "---":
            i += 1
            continue
        if s.startswith("|"):
            block = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                block.append(lines[i])
                i += 1
            rows = [_split_row(b) for b in block if not re.fullmatch(r"\|?\s*(:?-+:?\s*\|\s*)+(:?-+:?)?\s*\|?", b.strip())]
            if rows:
                _add_table(doc, rows)
            continue
        m = _IMG_MD.match(s)
        if m:
            p = (base_dir / m.group("src")).resolve()
            if p.exists():
                doc.add_picture(str(p), width=page_width)
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap = doc.add_paragraph(m.group("alt"), style="Caption")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_paragraph(f"[缺失图片：{m.group('src')}]")
            i += 1
            continue
        h = re.match(r"^(#{1,4})\s+(.*)$", s)
        if h:
            level = len(h.group(1))
            if level == 2 and not toc_done:
                doc.add_heading("目录", level=1)
                _add_toc_field(doc)
                doc.add_page_break()
                toc_done = True
            doc.add_heading(h.group(2), level=0 if level == 1 else level - 1)
            i += 1
            continue
        if s.startswith(">"):
            p = doc.add_paragraph(style="Quote")
            _add_inline(p, s.lstrip("> ").strip())
            i += 1
            continue
        if re.match(r"^[-*]\s+", s):
            p = doc.add_paragraph(style="List Bullet")
            _add_inline(p, re.sub(r"^[-*]\s+", "", s))
            i += 1
            continue
        if re.match(r"^\d+\.\s+", s):
            p = doc.add_paragraph(style="List Number")
            _add_inline(p, re.sub(r"^\d+\.\s+", "", s))
            i += 1
            continue
        # 普通段落：合并连续的非空、非特殊行
        buf = [s]
        i += 1
        while i < len(lines):
            nxt = lines[i].strip()
            if (not nxt or nxt.startswith(("|", "#", ">", "!["))
                    or re.match(r"^([-*]|\d+\.)\s+", nxt) or nxt == "---"):
                break
            buf.append(nxt)
            i += 1
        p = doc.add_paragraph()
        _add_inline(p, "".join(buf))
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))


def write_reports(report_md: str, md_path: str | Path, html_path: str | Path,
                  docx_path: str | Path | None = None, embed_images: bool = True) -> dict:
    """落盘 Markdown / HTML（/ Word）。返回实际写出的路径。"""
    md_path = Path(md_path)
    html_path = Path(html_path)
    md_path.parent.mkdir(parents=True, exist_ok=True)
    html_path.parent.mkdir(parents=True, exist_ok=True)
    md_path.write_text(report_md, encoding="utf-8")
    html_path.write_text(markdown_to_html(report_md, md_path.parent, embed_images), encoding="utf-8")
    out = {"md": md_path, "html": html_path}
    if docx_path is not None:
        markdown_to_docx(report_md, md_path.parent, Path(docx_path))
        out["docx"] = Path(docx_path)
    return out
