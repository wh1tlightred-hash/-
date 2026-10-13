"""扩展分析（交叉验证补充分析、扩展 SHAP、报告导出）的行为测试。

均在真实数据或小型合成数据上运行，不依赖 main.py 产物；预算调到很小以保证速度。
"""
import copy

import numpy as np
import pandas as pd
import pytest
from sklearn.model_selection import RepeatedStratifiedKFold, StratifiedKFold

from src.config import load_config
from src.cv_analysis import (
    calibration_summary,
    cv_scheme_comparison,
    learning_curves,
    multiclass_brier,
    permutation_tests,
    pooled_cv_predictions,
)
from src.data import load_dataset, split_xy
from src.eval import oof_metrics, refit_on_full_data, run_nested_cv
from src.explain import cv_shap, multi_model_shap, shap_values_for
from src.metrics import class_scores
from src.models import build_pipeline
from src.report_export import markdown_to_docx, markdown_to_html, write_reports


@pytest.fixture(scope="module")
def cfg():
    c = copy.deepcopy(load_config())
    c["analysis"]["cv_schemes"] = {"kfold": [3], "loo": True, "repeated": {"n_splits": 3, "n_repeats": 2}}
    c["analysis"]["permutation"] = {"n_permutations": 4, "cv_splits": 3}
    c["analysis"]["learning_curve"] = {"per_class_sizes": [2, 4, 9], "n_splits": 3, "n_repeats": 1}
    c["shap"]["nsamples_kernel"] = 40
    c["n_iter_random"] = 2
    return c


@pytest.fixture(scope="module")
def data(cfg):
    df = load_dataset(cfg["input_xlsx"], cfg["feature_columns"])
    X, y, meta = split_xy(df, cfg["feature_columns"])
    return np.asarray(X, dtype=float), y, meta


# ---------------------------------------------------------------- 交叉验证补充分析

def test_pooled_cv_covers_every_sample_once(data):
    X, y, _ = data
    r = pooled_cv_predictions("gnb", X, y, StratifiedKFold(3, shuffle=True, random_state=0), seed=0)
    assert r["y_pred"].shape == (27,) and np.all(r["y_pred"] >= 0)
    assert np.all(np.isfinite(r["y_score"]))
    assert len(r["per_fold"]) == 3


def test_pooled_cv_rejects_overlapping_splitter(data):
    """重复 CV 的测试折互相重叠，不能拼成「一次」折外预测，必须报错而不是悄悄覆盖。"""
    X, y, _ = data
    with pytest.raises(ValueError, match="重叠"):
        pooled_cv_predictions("gnb", X, y, RepeatedStratifiedKFold(n_splits=3, n_repeats=2, random_state=0), seed=0)


def test_loo_dummy_pathology_is_reported_not_hidden(cfg, data):
    """类别严格均衡时，先验基线在留一法下必然全错（Accuracy=0）——报告依赖这一已知性质。"""
    X, y, _ = data
    out = cv_scheme_comparison(X, y, cfg, seed=42, models=["dummy"])
    s = out["summary"].set_index("scheme")
    assert s.loc["loo", "accuracy"] == 0.0
    assert set(s.index) == {"skf3", "loo", "rskf3x2"}
    reps = out["units"][out["units"]["unit"] == "repeat"]
    assert len(reps) == 2


def test_learning_curve_subsamples_are_class_balanced(cfg, data):
    X, y, _ = data
    lc = learning_curves(X, y, cfg, seed=42, models=["gnb"])
    assert set(lc["n_train"]) <= {6, 12}            # 每类 9 超过训练折容量，被跳过
    assert (lc["n_train"] == 3 * lc["per_class"]).all()
    assert lc["test_f1_macro"].between(0, 1).all()


def test_permutation_pvalue_bounds(cfg, data):
    X, y, _ = data
    t = permutation_tests(X, y, cfg, seed=42, models=["gnb"])["table"].iloc[0]
    n = t["n_permutations"]
    assert 1 / (n + 1) - 1e-12 <= t["p_value"] <= 1.0


def test_multiclass_brier_reference_values():
    y = np.array([0, 1, 2])
    assert multiclass_brier(y, np.eye(3)) == pytest.approx(0.0)
    assert multiclass_brier(y, np.full((3, 3), 1 / 3)) == pytest.approx(2 / 3)


def test_calibration_summary_shapes(cfg, data):
    X, y, meta = data
    oof = {("gnb", "tuned"): pd.DataFrame({
        "y_true": y, "y_pred": y,
        **{f"score_{c}_{n}": (y == c).astype(float) * 0.8 + 0.1 * (y != c) for c, n in enumerate(cfg["class_names"])},
    })}
    cal = calibration_summary(oof, cfg, ["gnb"])
    row = cal["table"].iloc[0]
    assert row["accuracy"] == 1.0 and 0 <= row["brier"] < 0.1
    assert set(cal["curves"]["gnb"]) == {0, 1, 2}


# ---------------------------------------------------------------- 嵌套 CV 保留折模型

def test_keep_models_reproduces_oof_predictions(cfg, data):
    """保留下来的折模型必须正是产出 OOF 的那个模型：用它重新预测测试折，结果完全一致。"""
    X, y, meta = data
    res = run_nested_cv(pd.DataFrame(X), y, meta, cfg, seed=42, model_names=["gnb"], keep_models=True)
    fitted = res["fitted"][("gnb", "tuned")]
    assert len(fitted) == len(res["splits"]) == cfg["outer_cv"]["n_splits"]
    oof = res["oof"][("gnb", "tuned")]
    for fold_i, ((tr, te), est) in enumerate(zip(res["splits"], fitted)):
        got = est.predict(X[te])
        want = oof[oof["fold"] == fold_i]["y_pred"].to_numpy()
        np.testing.assert_array_equal(got, want)


def test_keep_models_off_by_default(cfg, data):
    X, y, meta = data
    res = run_nested_cv(pd.DataFrame(X), y, meta, cfg, seed=42, model_names=["dummy"])
    assert res["fitted"] == {}


# ---------------------------------------------------------------- SHAP 扩展

def test_tree_shap_is_additive_in_probability_space(data):
    """随机森林 TreeExplainer：基准值 + 各特征 SHAP 之和 = predict_proba（SHAP 的局部准确性）。"""
    X, y, _ = data
    pipe = build_pipeline("random_forest", seed=0, params={"n_estimators": 30}).fit(X, y)
    r = shap_values_for(pipe, X, X, {"shap": {}})
    recon = r["base_values"][None, :] + r["values"].sum(axis=1)
    np.testing.assert_allclose(recon, pipe.predict_proba(X), atol=1e-6)
    assert r["output_space"] == "predict_proba 概率"


def test_xgboost_shap_labelled_as_margin(data):
    X, y, _ = data
    pipe = build_pipeline("xgboost", seed=0, params={"n_estimators": 20}).fit(X, y)
    r = shap_values_for(pipe, X, X, {"shap": {}})
    assert "margin" in r["output_space"]
    margin = pipe.named_steps["clf"].predict(X, output_margin=True)
    np.testing.assert_allclose(r["base_values"][None, :] + r["values"].sum(axis=1), margin, atol=1e-4)


def test_cv_shap_covers_all_samples_with_out_of_fold_models(cfg, data):
    X, y, _ = data
    splits = list(StratifiedKFold(3, shuffle=True, random_state=0).split(X, y))
    fitted = [build_pipeline("random_forest", seed=0, params={"n_estimators": 20}).fit(X[tr], y[tr])
              for tr, _ in splits]
    out = cv_shap(fitted, splits, X, cfg["feature_columns"], cfg["class_names"], cfg)
    assert out["values"].shape == (27, 10, 3) and np.isfinite(out["values"]).all()
    fr = out["fold_ranks"]
    assert len(fr) == 3 * 10
    assert all(sorted(g["rank"]) == list(range(1, 11)) for _, g in fr.groupby("fold"))


def test_multi_model_shap_shares_sum_to_one(cfg, data):
    X, y, _ = data
    fitted = {m: build_pipeline(m, seed=0).fit(X, y) for m in ("gnb", "random_forest")}
    t = multi_model_shap(fitted, X, cfg["feature_columns"], cfg)["table"]
    for _, g in t.groupby("model"):
        assert g["share"].sum() == pytest.approx(1.0)
        assert sorted(g["rank"]) == list(range(1, 11))


def test_refit_on_full_data_matches_search_space(cfg, data):
    X, y, _ = data
    est, params, score, _ = refit_on_full_data(cfg, X, y, "gnb", seed=42)
    assert set(params) == {"var_smoothing"}
    assert 0 <= score <= 1
    assert class_scores(est, X).shape == (27, 3)


# ---------------------------------------------------------------- 报告导出

_MD = """# 标题

> 说明

## 1 第一章

正文 **加粗** 与 `代码`。

| 模型 | 值 |
| --- | --- |
| A | 0.5 |

- 列表项

![示例图](fig.png)
"""


@pytest.fixture()
def tiny_png(tmp_path):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt

    fig, ax = plt.subplots(figsize=(1, 1))
    ax.plot([0, 1], [0, 1])
    fig.savefig(tmp_path / "fig.png")
    plt.close(fig)
    return tmp_path


def test_html_embeds_images_and_builds_toc(tiny_png):
    html = markdown_to_html(_MD, tiny_png, embed_images=True)
    assert "data:image/png;base64," in html
    assert 'class="toc"' in html and "1 第一章" in html
    assert 'src="fig.png"' not in html
    assert 'src="fig.png"' in markdown_to_html(_MD, tiny_png, embed_images=False)


def test_docx_contains_table_image_and_headings(tiny_png):
    import docx

    out = tiny_png / "r.docx"
    markdown_to_docx(_MD, tiny_png, out)
    d = docx.Document(str(out))
    assert len(d.tables) == 1 and d.tables[0].cell(1, 1).text == "0.5"
    assert len(d.inline_shapes) == 1
    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading") or p.style.name == "Title"]
    assert "1 第一章" in heads and "目录" in heads


def test_write_reports_returns_all_formats(tiny_png):
    out = write_reports(_MD, tiny_png / "a.md", tiny_png / "a.html", tiny_png / "a.docx")
    assert set(out) == {"md", "html", "docx"} and all(p.exists() for p in out.values())
    assert (tiny_png / "a.md").read_text(encoding="utf-8") == _MD  # Markdown 原样保存（相对路径）
